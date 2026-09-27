import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_pricing_docx_report(relatorio_data: list) -> io.BytesIO:
    doc = Document()

    # Configuração de Estilo / Cabeçalho
    title = doc.add_heading("Relatório Comparativo de Preços", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_sub = doc.add_paragraph(
        "Análise de Margens, Concorrência e Sugestão de Preços - Motor de Preços 2.0"
    )
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # Espaçamento

    # Tabela Resumo
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    headers = [
        "SKU",
        "Produto",
        "Custo Base",
        "Menor Conc.",
        "Preço Sugerido",
        "Margem",
    ]
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text

    for item in relatorio_data:
        row_cells = table.add_row().cells
        row_cells[0].text = str(item.get("sku", ""))
        row_cells[1].text = str(item.get("nome", ""))
        row_cells[2].text = f"R$ {item.get('custo_base', 0):.2f}"

        menor = item.get("menor_concorrente")
        row_cells[3].text = f"R$ {menor:.2f}" if menor is not None else "N/A"

        row_cells[4].text = f"R$ {item.get('preco_sugerido', 0):.2f}"
        row_cells[5].text = f"{item.get('margem_minima', 0) * 100:.1f}%"

    doc.add_paragraph()
    doc.add_paragraph(
        "Relatório gerado automaticamente pelo sistema Motor de Preços 2.0."
    )

    target_stream = io.BytesIO()
    doc.save(target_stream)
    target_stream.seek(0)
    return target_stream
