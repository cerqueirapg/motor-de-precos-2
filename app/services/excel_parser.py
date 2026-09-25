import asyncio
from collections.abc import AsyncGenerator
from typing import Any

from openpyxl import load_workbook


async def parse_excel_streaming(
    file_path_or_stream, batch_size: int = 1000
) -> AsyncGenerator[list[dict[str, Any]], None]:
    """
    Lê uma planilha Excel em modo streaming (read_only=True) e gera
    lotes de dicionários para processamento eficiente em memória.
    """

    def _open_workbook():
        return load_workbook(
            filename=file_path_or_stream, read_only=True, data_only=True
        )

    workbook = await asyncio.to_thread(_open_workbook)
    sheet = workbook.active

    # Garante ao linter que a aba ativa não é nula
    if sheet is None:
        workbook.close()
        raise ValueError("A planilha enviada não possui uma aba ativa válida.")

    headers: list[str] = []
    batch: list[dict[str, Any]] = []

    for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
        if row_idx == 0:
            headers = [
                str(cell).strip() if cell is not None else f"column_{i}"
                for i, cell in enumerate(row)
            ]
            continue

        if not any(row):
            continue

        row_dict = {
            headers[i]: (row[i] if i < len(row) else None) for i in range(len(headers))
        }
        batch.append(row_dict)

        if len(batch) >= batch_size:
            yield batch
            batch = []

    if batch:
        yield batch

    workbook.close()
