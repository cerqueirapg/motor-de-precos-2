from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.schemas.product import PricingRequest, ProductBase
from app.services.calculator import PriceCalculatorService


def load_products_from_excel(file_path: str):
    """Lê a planilha Excel e agrupa os preços dos concorrentes por produto."""
    wb = load_workbook(filename=file_path, data_only=True)
    test_data = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) <= 1:
            continue

        # Extrai metadados do produto a partir da primeira linha de dados
        first_row = rows[1]
        product = ProductBase(
            sku=str(first_row[0]),
            name=str(first_row[1]),
            cost_price=Decimal(str(first_row[2])),
        )

        # Extrai todos os preços dos concorrentes (coluna index 3)
        competitor_prices = [
            Decimal(str(row[3])) for row in rows[1:] if row[3] is not None
        ]

        test_data.append((product, competitor_prices))

    return test_data


# Carrega a massa diretamente do Excel
EXCEL_TEST_DATA = load_products_from_excel("tests/fixtures/precos_concorrentes.xlsx")


@pytest.mark.parametrize("product, competitor_prices", EXCEL_TEST_DATA)
def test_calculate_price_from_excel(
    product: ProductBase, competitor_prices: list[Decimal]
):
    """Valida o cálculo com dados originados do arquivo .xlsx."""
    request = PricingRequest(
        product=product,
        desired_margin=Decimal("0.25"),
        min_margin=Decimal("0.10"),
        marketplace_tax=Decimal("0.12"),
        competitor_prices=competitor_prices,
    )

    response = PriceCalculatorService.calculate_price(request)

    assert response.sku == product.sku
    assert response.suggested_price > Decimal("0.00")
    assert response.adjusted_competitor_avg > Decimal("0.00")
    assert response.viability_status in [
        "excelente",
        "competitivo_ajustado",
        "incompetitivo",
    ]
